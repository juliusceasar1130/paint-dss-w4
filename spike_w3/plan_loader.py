# -*- coding: utf-8 -*-
"""周计划解析器（T1-lite；docs/06 v1.2 契约冻结 = §1，行级规则 = §2.1）。

分层纪律（同 T2/T4）：
- 纯函数层：parse_csv / parse_excel / load_plan / resolve_md_date / _parse_rows（无 DB）
- 编排层（T1.2 追加）：persist_plan / ensure_weekly_plan_tables（写 dss）
- 因果兜底（T1.3 追加）：infer_shift_causal / resolve_shift / recent_shift_history
- 坏行 = 跳过 + bad_rows 留痕，不崩溃、不当默认值；表级失败 = ParseError（不牵连他文件）
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Optional

from spike_w2 import shifts
from spike_w3.alert_engine import split_target

# 6 列规范（docs/08）；前 4 列缺一 = 表级 ParseError，后 2 列选填
CORE_COLS = ("日期", "日班工时", "中班工时", "出车计划")
ALL_COLS = ("日期", "日班工时", "中班工时", "出车计划", "进车计划", "备注")
REST_LIKE = ("", "0", "休")
VALID_HOURS = (8, 10, 11)
# 命名识别（T1.2）：year 可选；KW 后允许任意后缀（_sample/_b 等不排除）；扩展名 csv/xlsx/xls
PLAN_NAME_RE = re.compile(r"^week_plan_(\d{4})?KW(\d{2}).*\.(csv|xlsx|xls)$", re.IGNORECASE)
_DATE_FULL_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_MD_RE = re.compile(r"^(\d{1,2})/(\d{1,2})$")


class ParseError(Exception):
    """文件级解析失败（核心列表头缺失 / 结构损坏）——该文件留痕，不牵连其他文件。"""


@dataclass(frozen=True)
class PlanRow:
    date: date
    day_h: Optional[int]      # 8/10/11；None = 休/停产
    mid_h: Optional[int]
    total: Optional[int]      # 全天出车计划；None = 无计划
    entry: Optional[int]      # 进车计划（选填）
    note: str
    source: str               # 'file' | 'inferred'
    source_file: Optional[str]


@dataclass
class WeeklyPlan:
    rows: dict = field(default_factory=dict)
    bad_rows: list = field(default_factory=list)
    file_hashes: dict = field(default_factory=dict)

    # ------------------------------------------------------------------ #
    # 静态计划内容（agy 低 L2：在线运行路径必须走 resolve_shift，不得直接
    # 用 shift_windows 循环——缺日与双侧休同返 []，直接循环会静默绕过兜底）
    # ------------------------------------------------------------------ #
    def shift_windows(self, d: date) -> list:
        r = self.rows.get(d)
        if r is None:
            return []
        return shifts.shift_windows(
            d,
            day_type=f"{r.day_h}h" if r.day_h else None,
            mid_type=f"{r.mid_h}h" if r.mid_h else None,
        )

    def plan_target(self, d: date) -> tuple:
        """(日班目标, 中班目标)；复用 alert_engine.split_target（F10 单一事实源）。
        total None / 无行 → (None, None)——目标不猜（2026-09-27 用户拍板③）。"""
        r = self.rows.get(d)
        if r is None or r.total is None:
            return (None, None)
        return split_target(r.day_h or 0, r.mid_h or 0, r.total)

    def rest_days(self) -> set:
        """日+中皆休 = 整日停产（expected_silence，同 load_rest_days 语义）。"""
        return {d for d, r in self.rows.items() if r.day_h is None and r.mid_h is None}

    def missing(self, d0: date, d1: date) -> list:
        """[d0, d1] 闭区间内缺日清单（覆盖校验）。"""
        out, d = [], d0
        while d <= d1:
            if d not in self.rows:
                out.append(d)
            d += timedelta(days=1)
        return out


# ---------------------------------------------------------------------- #
# 行级解析（CSV/Excel 共享）
# ---------------------------------------------------------------------- #
def _parse_int(val) -> Optional[int]:
    """空 → None；int 直接；float 仅整数值（920.5 = 坏行，agy M2）；
    str 先 int(s)，失败再 int(float(s)) + 整数值校验（"920.0" → 920 合法）。"""
    if val is None:
        return None
    if isinstance(val, bool):
        raise ValueError(f"非法整数: {val!r}")
    if isinstance(val, int):
        return val
    if isinstance(val, float):
        if val.is_integer():
            return int(val)
        raise ValueError(f"非整数: {val!r}")
    s = str(val).strip()
    if s == "":
        return None
    try:
        return int(s)
    except ValueError:
        try:
            f = float(s)
        except ValueError:
            raise ValueError(f"非法整数: {val!r}") from None
        if f.is_integer():
            return int(f)
        raise ValueError(f"非整数: {val!r}")


def _parse_hours(val) -> Optional[int]:
    """F12 预校验：{8,10,11}（含 '11'/'11h'/Excel 11.0 整值）→ int；
    {休,0,空} → None；其他（'12' 笔误等）→ ValueError = 坏行。"""
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    if isinstance(val, int):
        if val in VALID_HOURS:
            return val
        if val == 0:
            return None
        raise ValueError(f"非法工时值: {val!r}")
    s = "" if val is None else str(val).strip().lower()
    if s in REST_LIKE:
        return None
    t = shifts.normalize_shift_type(s)
    if t in ("11h", "10h", "8h"):
        return int(t[:-1])
    raise ValueError(f"非法工时值: {val!r}")


def resolve_md_date(file_year: int, file_kw: int, month: int, day: int) -> Optional[date]:
    """M/D 候选年份算法（agy 二轮 M-new）：文件年 ±1 三候选，
    ISO (年,周) 与文件名 (年,KW) 匹配者胜出；全不中 → None（真坏行）。
    解唯一（ISO 周跨自然年的日期只落在 Y±1 中恰好一个年份，全扫描验证）。"""
    for cand_y in (file_year, file_year - 1, file_year + 1):
        try:
            d = date(cand_y, month, day)
        except ValueError:
            continue
        if d.isocalendar()[:2] == (file_year, file_kw):
            return d
    return None


def _file_week_meta(name: str) -> tuple:
    m = PLAN_NAME_RE.match(name)
    fy = int(m.group(1)) if m.group(1) else date.today().year
    return fy, int(m.group(2))


def _parse_date(val, file_year: int, file_kw: int) -> date:
    if isinstance(val, datetime):
        val = val.date()
    if isinstance(val, date):
        d = val
    else:
        s = str(val).strip()
        if _DATE_FULL_RE.match(s):
            d = date.fromisoformat(s)
        elif _MD_RE.match(s):
            mm, dd = _MD_RE.match(s).groups()
            d = resolve_md_date(file_year, file_kw, int(mm), int(dd))
            if d is None:
                raise ValueError(f"M/D 无法定位: {s}")
        else:
            raise ValueError(f"日期不可解析: {str(val)!r}")
    if d.isocalendar()[:2] != (file_year, file_kw):
        raise ValueError(f"ISO 周不符（文件 KW{file_kw:02d}）: {d}")
    return d


def _validate_header(cells) -> dict:
    """表头单元格 strip 后比对（agy C4-3）；4 核心列缺一 → ParseError。"""
    stripped = [str(c).strip() if c is not None else "" for c in cells]
    idx = {}
    for i, h in enumerate(stripped):
        if h and h not in idx:
            idx[h] = i
    for col in CORE_COLS:
        if col not in idx:
            raise ParseError(f"核心列缺失: {col}")
    return idx


def _rows_from_grid(grid: list, name: str) -> tuple:
    """grid = 含表头的二维表；返回 (rows, bad_rows)。空表 → 空 + 「文件为空」。"""
    if not grid or all(all(c is None or str(c).strip() == "" for c in r) for r in grid):
        return {}, [(name, "文件为空")]
    idx = _validate_header(grid[0])
    fy, kw = _file_week_meta(name)
    rows, bad = {}, []
    for raw in grid[1:]:
        if all(c is None or str(c).strip() == "" for c in raw):
            continue
        get = lambda col: raw[idx[col]] if (col in idx and idx[col] < len(raw)) else None
        try:
            d = _parse_date(get("日期"), fy, kw)
            day_h = _parse_hours(get("日班工时"))
            mid_h = _parse_hours(get("中班工时"))
            total = _parse_int(get("出车计划"))
            entry = _parse_int(get("进车计划"))
        except ValueError as e:
            bad.append((name, f"坏行: {e}"))
            continue
        note = str(get("备注") or "").strip()
        if d in rows:
            bad.append((name, f"日期重复（后行覆盖）: {d}"))
        rows[d] = PlanRow(d, day_h, mid_h, total, entry, note, "file", name)
    return rows, bad


def parse_csv(path) -> WeeklyPlan:
    path = Path(path)
    name = path.name
    raw = path.read_bytes()
    if not raw.strip():
        return WeeklyPlan(rows={}, bad_rows=[(name, "文件为空")],
                          file_hashes={name: hashlib.sha256(raw).hexdigest()})
    text = raw.decode("utf-8-sig")          # BOM 有无都兼容
    grid = [row for row in csv.reader(io.StringIO(text))]
    rows, bad = _rows_from_grid(grid, name)
    return WeeklyPlan(rows=rows, bad_rows=bad,
                      file_hashes={name: hashlib.sha256(raw).hexdigest()})


def parse_excel(path) -> WeeklyPlan:
    from openpyxl import load_workbook
    path = Path(path)
    name = path.name
    raw = path.read_bytes()
    if not raw.strip():
        return WeeklyPlan(rows={}, bad_rows=[(name, "文件为空")],
                          file_hashes={name: hashlib.sha256(raw).hexdigest()})
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]           # 首 sheet（F8）
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    rows, bad = _rows_from_grid(grid, name)
    return WeeklyPlan(rows=rows, bad_rows=bad,
                      file_hashes={name: hashlib.sha256(raw).hexdigest()})


def load_plan(folder) -> WeeklyPlan:
    """目录扫描 + 跨文件合并：
    - 命名识别（PLAN_NAME_RE）之外的文件忽略不报错
    - 同日期多文件值一致 → 合并（后处理文件的 source_file 归属，字典序确定）
    - 同日期值不同 → ValueError（入库前拦截，docs/02 契约）
    - 坏行 / 表级 ParseError / 文件损坏 / .xls 不支持 → 均留痕不崩溃（启动健壮性：
      plan_inbox 里任何坏文件不得拖垮在线主循环）
    """
    folder = Path(folder)
    merged, bad, hashes = {}, [], {}
    if folder.is_dir():
        files = sorted(p for p in folder.iterdir() if p.is_file() and PLAN_NAME_RE.match(p.name))
    else:
        files = [folder] if (folder.is_file() and PLAN_NAME_RE.match(folder.name)) else []
    for p in files:
        raw = p.read_bytes()
        hashes[p.name] = hashlib.sha256(raw).hexdigest()   # 识别即记哈希（.xls/坏文件也留痕可查）
        if p.suffix.lower() == ".xls":
            bad.append((p.name, "不支持格式（.xls 请另存为 .xlsx 或导出 CSV）"))
            continue
        parse = parse_csv if p.suffix.lower() == ".csv" else parse_excel
        try:
            w = parse(p)
        except ParseError as e:
            bad.append((p.name, f"解析失败: {e}"))
            continue
        except Exception as e:   # 文件损坏 / 编码坏 / openpyxl 异常——留痕不崩
            bad.append((p.name, f"解析失败: {type(e).__name__}: {e}"))
            continue
        bad.extend(w.bad_rows)
        for d, r in w.rows.items():
            if d in merged:
                old = merged[d]
                if (old.day_h, old.mid_h, old.total, old.entry) != (r.day_h, r.mid_h, r.total, r.entry):
                    raise ValueError(
                        f"跨文件冲突: {d} {p.name} vs {old.source_file}"
                        f"（{(old.day_h, old.mid_h, old.total, old.entry)} vs "
                        f"{(r.day_h, r.mid_h, r.total, r.entry)}）——请修正其中一个文件")
                merged[d] = r               # 值一致 → 合并（后文件归属）
            else:
                merged[d] = r
    return WeeklyPlan(rows=merged, bad_rows=bad, file_hashes=hashes)


# ---------------------------------------------------------------------- #
# 编排层（写 dss；只写两张新表，幂等 upsert——agy 中 M1）
# ---------------------------------------------------------------------- #
DDL_PLAN = """CREATE TABLE IF NOT EXISTS dss.dim_weekly_plan (
    plan_date   DATE PRIMARY KEY,
    day_hours   SMALLINT,
    mid_hours   SMALLINT,
    total_plan  INTEGER,
    entry_plan  INTEGER,
    note        TEXT,
    week_label  TEXT,
    source      TEXT NOT NULL,
    source_file TEXT,
    loaded_at   TIMESTAMPTZ NOT NULL DEFAULT now()
)"""

DDL_META = """CREATE TABLE IF NOT EXISTS dss.work_plan_file_meta (
    file_name   TEXT PRIMARY KEY,
    sha256      TEXT NOT NULL,
    loaded_at   TIMESTAMPTZ NOT NULL DEFAULT now()
)"""


def ensure_weekly_plan_tables(conn) -> None:
    """懒建 DDL（幂等；只写 dss）。boot_sequence 中必须在 check_dim_tables 之前调（agy L3）。"""
    cur = conn.cursor()
    cur.execute(DDL_PLAN)
    cur.execute(DDL_META)
    conn.commit()


_UPSERT_PLAN = """INSERT INTO dss.dim_weekly_plan
    (plan_date, day_hours, mid_hours, total_plan, entry_plan, note,
     week_label, source, source_file, loaded_at)
VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, now())
ON CONFLICT (plan_date) DO UPDATE SET
    day_hours   = EXCLUDED.day_hours,   mid_hours   = EXCLUDED.mid_hours,
    total_plan  = EXCLUDED.total_plan,  entry_plan  = EXCLUDED.entry_plan,
    note        = EXCLUDED.note,        week_label  = EXCLUDED.week_label,
    source      = EXCLUDED.source,      source_file = EXCLUDED.source_file,
    loaded_at   = now()"""

_UPSERT_META = """INSERT INTO dss.work_plan_file_meta (file_name, sha256, loaded_at)
VALUES (%s, %s, now())
ON CONFLICT (file_name) DO UPDATE SET sha256 = EXCLUDED.sha256, loaded_at = now()"""


def persist_plan(conn, plan: WeeklyPlan) -> int:
    """文件哈希幂等入库（F/T1.2）：
    - 逐文件 SELECT meta 哈希：相同 → 跳过（0 写入）；不同/新文件 → upsert 该文件行 + meta
    - upsert 语义（agy M1）：文件改名安全（新名覆盖旧名写的行）；无 DELETE
    - 同事务；失败回滚保旧数据（影子纪律）；返回新写入行数
    """
    cur = conn.cursor()
    written = 0
    try:
        for name in sorted(plan.file_hashes):
            h = plan.file_hashes[name]
            cur.execute(
                "SELECT sha256 FROM dss.work_plan_file_meta WHERE file_name = %s", (name,))
            row = cur.fetchone()
            if row is not None and row[0] == h:
                continue                       # 同内容 → 跳过
            for d in sorted(d for d, r in plan.rows.items() if r.source_file == name):
                r = plan.rows[d]
                cur.execute(_UPSERT_PLAN, (
                    d, r.day_h, r.mid_h, r.total, r.entry, r.note,
                    shifts.iso_week_label(d), r.source, name))
                written += 1
            cur.execute(_UPSERT_META, (name, h))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return written


# ---------------------------------------------------------------------- #
# 因果兜底（T1.3）：只用已完日、只用 < now 的事件（因果红线）。
# 禁用 W2 detect_*（回放后视语义）——阈值常量复用 shifts，判定逻辑内联重写。
# ---------------------------------------------------------------------- #
_DAY_BUCKET_TIMES = (time(5, 0), time(6, 0), time(8, 0))
_DAY_BUCKET_HOURS = {time(5, 0): 11, time(6, 0): 10, time(8, 0): 8}


def classify_completed_day(day_evts, mid_exit_evts) -> tuple:
    """已完日班型事后判定 → (day_h, mid_h)；(None, None) = 休（零下线零事件）。

    day_evts: 当日产品车事件时刻（全日；15min 桶判 05:00/06:00/08:00 开班，
              阈值 = shifts.SURGE_MIN_15MIN / MIN_DAY_EVENTS，同 W2 口径）
    mid_exit_evts: 中班段 [16:30, 次 04:45) 内 4 下线点事件时刻
                   （末事件距 16:30 分钟数：≥630 → 11 / ≥510 → 10 / 其余 → 8；
                     < shifts.MID_MIN_FLOOR(300) = 数据截断 → None）
    """
    day_h = None
    if len(day_evts) >= shifts.MIN_DAY_EVENTS:
        d0 = day_evts[0].date()
        hit = None
        for tt in _DAY_BUCKET_TIMES:
            t0 = datetime.combine(d0, tt, tzinfo=day_evts[0].tzinfo)
            if sum(1 for t in day_evts if t0 <= t < t0 + timedelta(minutes=15)) >= shifts.SURGE_MIN_15MIN:
                hit = tt
                break
        day_h = _DAY_BUCKET_HOURS[hit] if hit is not None else 10   # 渐进爬坡回退（同 W2）
    if not mid_exit_evts:
        return (day_h, None) if day_evts else (None, None)
    last = max(mid_exit_evts)
    d_attr = last.date() - timedelta(days=1) if last.time() < time(4, 45) else last.date()
    mid_min = (last - datetime.combine(d_attr, time(16, 30), tzinfo=last.tzinfo)).total_seconds() / 60
    if mid_min < shifts.MID_MIN_FLOOR:
        return (day_h, None)
    mid_h = 11 if mid_min >= 630 else (10 if mid_min >= 510 else 8)
    return (day_h, mid_h)


def infer_shift_causal(history, weekday: int) -> tuple:
    """纯函数（无 DB）：同 weekday 样本逐侧众数（休 = None 不进侧投票）；
    平票 → 较新日值；无样本 → (11, 11)；有样本但双侧皆休 → (None, None)。
    返回 (day_h, mid_h, 'inferred')。14 天窗由 recent_shift_history 控制。"""
    samples = [(d, dh, mh) for (d, dh, mh) in history if d.weekday() == weekday]
    if not samples:
        return (11, 11, "inferred")

    def _mode(vals_d):
        if not vals_d:
            return None
        cnt = {}
        latest = {}
        for d, v in vals_d:
            cnt[v] = cnt.get(v, 0) + 1
            if v not in latest or d > latest[v]:
                latest[v] = d
        best = max(cnt.values())
        tied = [v for v, c in cnt.items() if c == best]
        return max(tied, key=lambda v: latest[v])          # 平票 → 较新日值

    day_m = _mode([(d, dh) for d, dh, mh in samples if dh is not None])
    mid_m = _mode([(d, mh) for d, dh, mh in samples if mh is not None])
    if day_m is None and mid_m is None:
        return (None, None, "inferred")                    # 有样本但全休
    return (day_m if day_m is not None else 11,
            mid_m if mid_m is not None else 11, "inferred")


def recent_shift_history(conn, now: datetime, weeks: int = 2) -> list:
    """因果 DB 构建器：[now-14d, now) 已完日 → [(date, day_h|None, mid_h|None)]。
    - 已完日 = 中班窗（≤次 04:45 收）已闭合 → 上界 = min(now.date()-1, max_date)
      （agy M4：max_date 之后 = 未观测，不入样本、不当休）
    - SET TIME ZONE；只读 ods.carbody_history；DATE_EVT < now 因果上界钉死
    - 禁用 W2 detect_*（回放后视语义）
    """
    from spike_w3.common import EXIT_STATIONS, CST
    cur = conn.cursor()
    cur.execute("SET TIME ZONE 'Asia/Shanghai'")
    cur.execute('SELECT max("DATE_EVT")::date FROM ods.carbody_history')
    row = cur.fetchone()
    max_date = row[0] if row and row[0] is not None else now.date() - timedelta(days=1)
    upper = min(now.date() - timedelta(days=1), max_date)
    span = 7 * weeks
    lower = upper - timedelta(days=span - 1)
    lo_dt = datetime.combine(lower, time(0, 0), tzinfo=CST)
    cur.execute(
        'SELECT "DATE_EVT" FROM ods.carbody_history '
        'WHERE "BODY_ID" LIKE \'78202%%\' AND "DATE_EVT" >= %s AND "DATE_EVT" < %s',
        (lo_dt, now))
    day_times = [r[0] for r in cur.fetchall()]
    exits = ", ".join(f"'{s}'" for s in EXIT_STATIONS)
    cur.execute(
        f'SELECT "DATE_EVT" FROM ods.carbody_history '
        f'WHERE "RW_STATION_ID" IN ({exits}) AND "DATE_EVT" >= %s AND "DATE_EVT" < %s',
        (lo_dt, now))
    exit_times = [r[0] for r in cur.fetchall()]
    day_c, mid_c = {}, {}
    for t in day_times:
        day_c.setdefault(t.date(), []).append(t)
    for t in exit_times:
        tm = t.time()
        if tm >= time(16, 30):
            mid_c.setdefault(t.date(), []).append(t)
        elif tm < time(4, 45):
            mid_c.setdefault(t.date() - timedelta(days=1), []).append(t)
        # 04:45–16:30 无归属微隙：弃
    out = []
    d = lower
    while d <= upper:
        out.append((d,) + classify_completed_day(day_c.get(d, []), mid_c.get(d, [])))
        d += timedelta(days=1)
    return out


def resolve_shift(d: date, plan: WeeklyPlan, history: list) -> tuple:
    """docs/02 冻结契约编排层：plan 有行 → 计划值 + 'plan'；无行 → 因果推断 + 'inferred'。"""
    r = plan.rows.get(d)
    if r is not None:
        return (r.day_h, r.mid_h, "plan")
    return infer_shift_causal(history, d.weekday())
